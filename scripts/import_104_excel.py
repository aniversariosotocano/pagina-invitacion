import os
import sqlite3
import json
import openpyxl

EXCEL_PATH = r"C:\Users\User\Downloads\Invitados_38_Aniversario_Soto_Cano_104.xlsx"
DB_PATH = os.path.join("assets", "data", "protocolo.db")
JSON_INVITADOS_PATH = os.path.join("assets", "data", "invitados.json")
JSON_CONFIG_PATH = os.path.join("assets", "data", "config.json")

def init_db(conn):
    cursor = conn.cursor()
    
    # 1. Tabla de Configuración Global del Protocolo
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS config_global (
        id INTEGER PRIMARY KEY CHECK (id = 1),
        fecha_fundacion TEXT NOT NULL DEFAULT '1988-09-01',
        aniversario INTEGER NOT NULL DEFAULT 38,
        aniversario_auto INTEGER NOT NULL DEFAULT 1,
        fecha_evento TEXT NOT NULL DEFAULT '24 de Septiembre de 2026',
        hora_evento TEXT NOT NULL DEFAULT '10:00 am',
        vestimenta TEXT NOT NULL DEFAULT 'Militar Uniforme D (Kepi) / Invitados Especiales Formal',
        updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # 2. Tabla de Invitados Oficiales
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS invitados (
        id TEXT PRIMARY KEY,
        no INTEGER NOT NULL,
        categoria TEXT NOT NULL,
        tratamiento TEXT NOT NULL,
        grado TEXT NOT NULL,
        nombre TEXT NOT NULL,
        cargo TEXT NOT NULL,
        activo INTEGER NOT NULL DEFAULT 0,
        observaciones TEXT,
        created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
        updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # 3. Índices Optimizados
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_invitados_activo ON invitados(activo);")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_invitados_no ON invitados(no);")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_invitados_categoria ON invitados(categoria);")
    cursor.execute("CREATE INDEX IF NOT EXISTS idx_invitados_nombre ON invitados(nombre);")

    # Inserción o actualización de la configuración por defecto
    cursor.execute("""
    INSERT INTO config_global (id, fecha_fundacion, aniversario, aniversario_auto, fecha_evento, hora_evento, vestimenta)
    VALUES (1, '1988-09-01', 38, 1, '24 de Septiembre de 2026', '10:00 am', 'Militar Uniforme D (Kepi) / Invitados Especiales Formal')
    ON CONFLICT(id) DO UPDATE SET
        fecha_fundacion=excluded.fecha_fundacion,
        aniversario=excluded.aniversario,
        aniversario_auto=excluded.aniversario_auto,
        fecha_evento=excluded.fecha_evento,
        hora_evento=excluded.hora_evento,
        vestimenta=excluded.vestimenta,
        updated_at=CURRENT_TIMESTAMP;
    """)
    conn.commit()

def determine_treatment(no, grado, cargo, nombre, grupo):
    if no == 1:
        return "Al: Sr. Comandante General De La FAH"
    if no == 2:
        return "Al: Sr. Jefe del Estado Mayor General Aéreo"
    if no == 3:
        return "Al: Sr. Inspector General de la FAH"
        
    g_lower = (grado or "").lower()
    c_lower = (cargo or "").lower()
    grp_lower = (grupo or "").lower()
    
    # Tratamiento femenino si corresponde
    female_markers = ['jefa', 'señora', 'doña', 'dra.', 'licda.', 'ing. ana', 'ana patricia', 'bethy', 'yilian', 'sidia', 'dulce', 'dolores', 'esposa', 'hija']
    is_female = any(m in (nombre or "").lower() or m in c_lower or m in g_lower for m in female_markers)
    
    if "familia" in grp_lower:
        return "A la distinguida:" if is_female else "Al distinguido:"
        
    if "civil" in grp_lower or "especiales civiles" in grp_lower:
        if "dr." in g_lower or "doctor" in g_lower:
            return "Al distinguido: Dr."
        if "ing." in g_lower or "ingeniero" in g_lower:
            return "Al distinguido: Ing."
        if "lic." in g_lower or "licenciado" in g_lower:
            return "Al distinguido: Lic."
        return "A la distinguida:" if is_female else "Al distinguido:"
        
    # Militares
    return "A la:" if is_female else "Al: Sr."

def import_excel():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb['Invitados']
    
    conn = sqlite3.connect(DB_PATH)
    init_db(conn)
    cursor = conn.cursor()
    
    # Limpiar tabla invitados para recarga fresca y precisa
    cursor.execute("DELETE FROM invitados;")
    
    invitados_list = []
    
    for r in range(5, ws.max_row + 1):
        val_no = ws.cell(r, 1).value
        if val_no is None:
            continue
        try:
            no = int(val_no)
        except ValueError:
            continue
            
        grupo = str(ws.cell(r, 2).value or "").strip()
        nombre = str(ws.cell(r, 3).value or "").strip()
        grado = str(ws.cell(r, 4).value or "").strip()
        cargo = str(ws.cell(r, 5).value or "").strip()
        observaciones = str(ws.cell(r, 7).value or "").strip()
        
        # Activo: Los primeros 5 activos, del 6 al 104 desactivados
        activo = 1 if no <= 5 else 0
        
        tratamiento = determine_treatment(no, grado, cargo, nombre, grupo)
        guest_id = f"invitado-{no}"
        
        cursor.execute("""
        INSERT INTO invitados (id, no, categoria, tratamiento, grado, nombre, cargo, activo, observaciones)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (guest_id, no, grupo, tratamiento, grado, nombre, cargo, activo, observaciones))
        
        invitados_list.append({
            "id": guest_id,
            "no": no,
            "categoria": grupo,
            "tratamiento": tratamiento,
            "grado": grado,
            "nombre": nombre,
            "cargo": cargo,
            "activo": bool(activo),
            "observaciones": observaciones
        })
        
    conn.commit()
    
    # Obtener configuración global
    cursor.execute("SELECT fecha_fundacion, aniversario, aniversario_auto, fecha_evento, hora_evento, vestimenta FROM config_global WHERE id = 1")
    cfg_row = cursor.fetchone()
    config_dict = {
        "fecha_fundacion": cfg_row[0],
        "aniversario": cfg_row[1],
        "aniversario_auto": bool(cfg_row[2]),
        "fecha_evento": cfg_row[3],
        "hora_evento": cfg_row[4],
        "vestimenta": cfg_row[5]
    }
    
    # Exportar JSON de sincronización estática
    with open(JSON_INVITADOS_PATH, 'w', encoding='utf-8') as f:
        json.dump(invitados_list, f, ensure_ascii=False, indent=2)
        
    with open(JSON_CONFIG_PATH, 'w', encoding='utf-8') as f:
        json.dump(config_dict, f, ensure_ascii=False, indent=2)
        
    # Validaciones en consola
    cursor.execute("SELECT COUNT(*) FROM invitados")
    total_count = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM invitados WHERE activo = 1")
    activos_count = cursor.fetchone()[0]
    cursor.execute("SELECT COUNT(*) FROM invitados WHERE activo = 0")
    inactivos_count = cursor.fetchone()[0]
    
    print(f"=== MIGRACIÓN A SQLITE COMPLETADA CON ÉXITO ===")
    print(f"Base de datos SQLite: {DB_PATH}")
    print(f"Total registros cargados: {total_count}")
    print(f"Registros Activos (Habilitados): {activos_count} (Primeros 5)")
    print(f"Registros Inactivos (Desactivados): {inactivos_count}")
    print(f"Índices optimizados creados: idx_invitados_activo, idx_invitados_no, idx_invitados_categoria, idx_invitados_nombre")
    print(f"Archivos JSON sincronizados: {JSON_INVITADOS_PATH}, {JSON_CONFIG_PATH}")
    
    conn.close()

if __name__ == '__main__':
    import_excel()
